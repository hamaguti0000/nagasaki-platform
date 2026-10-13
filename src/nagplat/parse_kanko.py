"""県観光統計 集約表(観光客数 平成21年-)の解析。

シート = 指標。行 = 県計・ブロック別・21市町、列 = 年。
赤字(FFFF0000)のセルは、県が「再算定後」の値として朱書きしたもの(シート末尾の注記: 「※朱書きの年は、再算定後」)。
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from .common import AREA_NAMES, PREF_TOTAL

# シート名 → 列名
SHEETS = {
    "観光客延べ数": "tourists_total",
    "日帰り客数": "day_visitors",
    "宿泊客延べ滞在数": "stay_person_days",
    "延べ宿泊客数": "stay_nights",
    "宿泊客実数": "stay_guests",
    "観光客実数": "visitors_actual",
}
RED = "FFFF0000"
INDEXED_RED = 10  # Excelの標準パレットのインデックス10 = 赤。7セルがこの指定(rgb指定の赤と混在する)


def _is_red(cell) -> bool:
    c = cell.font.color
    if c is None:
        return False
    return (c.type == "rgb" and c.rgb == RED) or (c.type == "indexed" and c.indexed == INDEXED_RED)


def parse_aggregate(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path)  # data_only は不要(値のセル)
    out: dict[tuple[int, str], dict] = {}
    for sheet, col in SHEETS.items():
        ws = wb[sheet]
        header = next(r for r in ws.iter_rows(min_row=1, max_row=6) if r[0].value == "市町名")
        years = {}
        for c in header:
            m = re.search(r"\((\d{4})\)", str(c.value or ""))
            if m:
                years[c.column - 1] = int(m.group(1))
        for row in ws.iter_rows(min_row=header[0].row + 1):
            name = (row[0].value or "").strip() if isinstance(row[0].value, str) else ""
            if name not in AREA_NAMES and name != PREF_TOTAL:
                continue
            for idx, y in years.items():
                cell = row[idx]
                v = cell.value
                d = out.setdefault((y, name), {"year": y, "area_name": name})
                d[col] = float(v) if isinstance(v, (int, float)) else None
                d[f"{col}__recalc"] = _is_red(cell)
    return [out[k] for k in sorted(out)]
